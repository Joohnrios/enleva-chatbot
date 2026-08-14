"""Completa a planilha de custos: faixa $2/$10 vs $3/$15 + 5 intensidades (base 350)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SRC = Path(r"C:\Users\John-Rios\Documents\enleva-chatbot\docs\estimativa-custo-api-bot-enleva.xlsx")
SRC_ALT = Path(
    r"C:\Users\John-Rios\Documents\enleva-chatbot\docs\estimativa-custo-api-bot-enleva-v2.xlsx"
)
DOWNLOADS = Path(r"c:\Users\John-Rios\Downloads\estimativa-custo-api-bot-enleva.xlsx")

YELLOW = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def main() -> None:
    wb = load_workbook(SRC)
    pre = wb["Premissas"]

    # --- Preços: faixa (vigente + estresse) ---
    pre["A8"] = "Preços da API Anthropic (USD por milhão de tokens) — FAIXA para gestores"
    pre["A9"] = "Modelo / cenário"
    pre["B9"] = "Entrada"
    pre["C9"] = "Saída"
    pre["D9"] = "Nota"
    for col in ("A", "B", "C", "D"):
        pre[f"{col}9"].font = BOLD

    pre["A10"] = "Sonnet 5 — vigente / permanente (usar como base)"
    pre["B10"] = 2
    pre["C10"] = 10
    pre["D10"] = (
        "Preço oficial atual (consultado ago/2026). Em 10/08/2026 a Anthropic anunciou que "
        "$2/$10 deixou de ser só promocional até 31/08 e passou a ser o preço padrão; "
        "o aumento para $3/$15 em 01/09/2026 foi cancelado. Fonte: platform.claude.com/docs pricing."
    )
    pre["B10"].fill = YELLOW
    pre["C10"].fill = YELLOW

    pre["A11"] = "Sonnet 5 — cenário ESTRESSE $3/$15 (não é o list price vigente)"
    pre["B11"] = 3
    pre["C11"] = 15
    pre["D11"] = (
        "Mantido na planilha só como sensibilidade conservadora (se o list price mudar no futuro "
        "ou para comparar com material antigo). NÃO apresentar como 'preço pós-31/08' — isso "
        "estava previsto e foi cancelado. Haiku permanece $1/$5 nas linhas abaixo."
    )
    pre["B11"].fill = YELLOW
    pre["C11"].fill = YELLOW

    # Shift Haiku - was A11, now need A12. Check if A12 exists for tokens section
    # Current structure after our earlier edit: A11 was Haiku. We're overwriting.
    # Put Haiku at A12 and shift token section notes - read what A12+ was

    pre["A12"] = "Claude Haiku 4.5 (classificador de escopo — stack atual)"
    pre["B12"] = 1
    pre["C12"] = 5
    pre["D12"] = (
        "No stack atual Haiku só classifica escopo ambíguo. Fonte: platform.claude.com — pricing."
    )

    # Token profile was at A13+ - may still be there. Update Cálculo formulas that pointed to B10/C10/B11/C11
    # Old: Premissas B10/C10 = Sonnet, B11/C11 = Haiku
    # New: B10/C10 = Sonnet vigente, B11/C11 = Sonnet stress, B12/C12 = Haiku

    calc = wb["Cálculo"]
    # Update Haiku refs from B11/C11 to B12/C12
    calc["B13"] = (
        "=(Premissas!$B$25/1000000*Premissas!$B$12)"
        "+(Premissas!$B$26/1000000*Premissas!$C$12)"
    )
    # Sonnet generation still uses B10/C10 (vigente)
    # Add parallel columns for stress pricing on Amigável

    # Extend Cálculo with stress price column
    calc["D5"] = "Amigável @ $3/$15 (estresse)"
    calc["D5"].font = BOLD
    calc["D6"] = "=C6"
    calc["D7"] = "=C7"
    calc["D8"] = "=D6/1000000*Premissas!$B$11"
    calc["D9"] = "=D7/1000000*Premissas!$C$11"
    calc["D10"] = "=D8+D9"
    calc["D17"] = "Amigável @ $3/$15"
    calc["D18"] = "=D10+$B$14"
    calc["D19"] = "=D18*Premissas!$B$6"

    # Labels
    calc["A10"] = "Custo Sonnet por mensagem (USD) — geração"
    calc["C5"] = "Amigável @ $2/$10 (vigente)"

    # --- Nova aba: Intensidade (pico 350) ---
    if "Intensidade 350" in wb.sheetnames:
        del wb["Intensidade 350"]
    inten = wb.create_sheet("Intensidade 350", 3)

    inten["A1"] = "Cenários de intensidade — pico de usuários ativos no bot"
    inten["A1"].font = Font(bold=True, size=14)
    inten["A2"] = (
        "Origem do '350': estimativa informada pelo solicitante na discussão de custo "
        "(não veio de RH/TI formal nem inventada pelo modelo). Célula amarela = editável — "
        "confirmar com RH/TI antes de cravar na apresentação aos gestores."
    )
    inten["A2"].alignment = WRAP
    inten.row_dimensions[2].height = 48

    inten["A4"] = "Usuários ativos no bot (pico / mês)"
    inten["B4"] = 350
    inten["B4"].fill = YELLOW
    inten["D4"] = "Empresa total (referência Premissas)"
    inten["E4"] = "=Premissas!B5"

    inten["A5"] = "Câmbio USD→BRL"
    inten["B5"] = "=Premissas!B6"
    inten["A6"] = "Dias úteis/mês (cenários intensos)"
    inten["B6"] = 22
    inten["B6"].fill = YELLOW

    inten["A8"] = "Nível"
    inten["B8"] = "Como calcula msgs/mês"
    inten["C8"] = "Msgs/mês"
    inten["D8"] = "R$/mês @ $2/$10"
    inten["E8"] = "R$/mês @ $3/$15"
    inten["F8"] = "USD/mês @ $2/$10"
    inten["G8"] = "USD/mês @ $3/$15"
    for col in range(1, 8):
        inten.cell(8, col).font = BOLD

    # Cost per msg Amigável from Cálculo: C18 = $2/$10 stack, D18 = $3/$15 stack
    # Rows 9-13: five levels
    levels = [
        ("Leve", "usuários × 2 conversas × 2,5 msgs", "=$B$4*2*2.5"),
        ("Moderada", "usuários × 3 conversas × 3 msgs", "=$B$4*3*3"),
        ("Forte", "usuários × 4 conversas × 3,5 msgs", "=$B$4*4*3.5"),
        ("Intensa", "1 msg/usuário/dia útil", "=$B$4*$B$6*1"),
        ("Muito intensa", "2 msgs/usuário/dia útil", "=$B$4*$B$6*2"),
    ]
    for i, (name, how, msg_f) in enumerate(levels, start=9):
        inten[f"A{i}"] = name
        inten[f"B{i}"] = how
        inten[f"C{i}"] = msg_f
        inten[f"D{i}"] = f"=C{i}*Cálculo!$C$19"
        inten[f"E{i}"] = f"=C{i}*Cálculo!$D$19"
        inten[f"F{i}"] = f"=C{i}*Cálculo!$C$18"
        inten[f"G{i}"] = f"=C{i}*Cálculo!$D$18"

    inten["A15"] = "Leitura para gestores"
    inten["A15"].font = BOLD
    inten["A16"] = (
        "•  Apresentar sempre FAIXA: coluna $2/$10 (preço vigente/permanente segundo Anthropic) "
        "e coluna $3/$15 (estresse). Não diga que em 01/09 o preço 'sobe' — o aumento previsto "
        "foi cancelado em 10/08/2026; a coluna estresse é só prudência orçamentária."
    )
    inten["A17"] = (
        "•  Piloto real (pós CG-SI) provavelmente roda sob o list price vigente na data do go-live. "
        "Revalidar platform.claude.com/docs pricing na semana do contrato."
    )
    inten["A18"] = (
        "•  Se o pico de 350 for confirmado por RH/TI, o nível Moderada/Forte é o mais útil para "
        "orçamento; Intensa/Muito intensa só se o bot virar hábito diário da maioria."
    )
    inten["A19"] = (
        "•  Recomendação de produto: piloto em Anthropic Sonnet; teto/alerta de gasto; se apertar, "
        "bakeoff Haiku-como-gerador (mesmo fornecedor) — não Gemini no caminho crítico do piloto."
    )
    for r in range(16, 20):
        inten[f"A{r}"].alignment = WRAP
        inten.row_dimensions[r].height = 40

    inten.column_dimensions["A"].width = 18
    inten.column_dimensions["B"].width = 36
    inten.column_dimensions["C"].width = 14
    for col in ("D", "E", "F", "G"):
        inten.column_dimensions[col].width = 18

    # --- Resumo ---
    res = wb["Resumo"]
    res["A4"] = (
        "Apresentação: ver aba 'Intensidade 350' (faixa $2/$10 e $3/$15). "
        "Abaixo = Moderado clássico 600 colab. @ preço vigente $2/$10."
    )
    res["A11"] = (
        "•  Preço Sonnet 5: vigente $2/$10 (Anthropic confirmou permanente em 10/08/2026; "
        "aumento a $3/$15 em 01/09 foi cancelado). A planilha ainda mostra $3/$15 como "
        "coluna de ESTRESSE na aba Intensidade 350 — para gestores verem a faixa, não porque "
        "o piloto 'vai pagar promo até 31/08'."
    )
    res["A11"].alignment = WRAP
    res.row_dimensions[11].height = 55

    # Fix Comparativo Haiku price refs if they use Premissas B11 for Haiku
    comp = wb["Comparativo Provedores"]
    # B12 was Sonnet using Premissas B10 - OK
    # B13 Haiku as gen used Premissas B11 - NOW B11 is stress Sonnet, must use B12/C12
    if comp["B13"].value and "Premissas!$B$11" in str(comp["B13"].value):
        tin = "(Premissas!$B$14+Premissas!$B$15+Premissas!$B$16+Premissas!$B$17)"
        tout = "Premissas!$B$21"
        comp["B13"] = f"={tin}/1000000*Premissas!$B$12+{tout}/1000000*Premissas!$C$12"

    # Benefício esperado note
    ben = wb["Benefício esperado"]
    ben["A10"] = (
        "Compare o custo de API (abas Resumo / Intensidade 350 — faixa $2/$10 e estresse $3/$15) "
        "com o custo de oportunidade de RH/TI. Não invente ROI por chamado. "
        "Confirmar o número de usuários ativos (célula B4 da Intensidade 350) com RH/TI."
    )
    ben["A10"].alignment = WRAP

    # Print expected numbers for 350
    tin, tout = 1930, 220
    fx = 5.3
    haiku_w = (150 / 1e6 * 1 + 10 / 1e6 * 5) * 0.2
    for label, pin, pout in [("$2/$10", 2, 10), ("$3/$15", 3, 15)]:
        c = tin / 1e6 * pin + tout / 1e6 * pout + haiku_w
        print(label, "per msg USD", round(c, 5), "BRL", round(c * fx, 5))
        for name, msgs in [
            ("Leve", 350 * 2 * 2.5),
            ("Moderada", 350 * 3 * 3),
            ("Forte", 350 * 4 * 3.5),
            ("Intensa", 350 * 22),
            ("Muito intensa", 350 * 22 * 2),
        ]:
            print(f"  {name}: {int(msgs)} msgs → R$ {msgs * c * fx:.2f}")

    saved = []
    for path in (SRC, SRC_ALT, DOWNLOADS):
        try:
            wb.save(path)
            saved.append(str(path))
        except PermissionError:
            print("locked, skip:", path)
    print("saved:", saved)


if __name__ == "__main__":
    main()
