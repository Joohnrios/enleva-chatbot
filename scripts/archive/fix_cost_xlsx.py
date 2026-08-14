"""Corrige estimativa-custo-api-bot-enleva.xlsx (Sonnet $2/$10, OpenAI, textos)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SRC = Path(r"c:\Users\John-Rios\Downloads\estimativa-custo-api-bot-enleva.xlsx")
DST_DOCS = Path(r"C:\Users\John-Rios\Documents\enleva-chatbot\docs\estimativa-custo-api-bot-enleva.xlsx")


def main() -> None:
    wb = load_workbook(SRC)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    red_font = Font(color="C00000", bold=True)
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    pre = wb["Premissas"]
    pre["B10"] = 2
    pre["C10"] = 10
    pre["D10"] = (
        "Preço padrão vigente (ago/2026): $2,00 input / $10,00 output — o valor "
        "introdutivo tornou-se permanente (Anthropic não elevou para $3/$15 em 01/09/2026). "
        "Fonte: platform.claude.com/docs/en/about-claude/pricing"
    )
    pre["D11"] = (
        "Haiku no stack ATUAL = só classificador de escopo ambíguo "
        "(não gera a resposta RAG). Fonte: platform.claude.com — pricing."
    )

    res = wb["Resumo"]
    res["A10"] = (
        "•  O custo de API no cenário Moderado fica na casa de dezenas de R$/mês. No Agressivo, "
        "Sonnet vs a opção mais barata (Gemini Flash-Lite) pode chegar a ~R$ 120–185/mês de diferença — "
        "isso é material na linha de API, mas ainda pequeno frente a 1 FTE de RH/TI. Há tempo para bakeoff "
        "de qualidade antes de trocar provedor."
    )
    res["A11"] = (
        "•  O preço do Claude Sonnet 5 usado aqui é o padrão vigente $2/$10 por milhão de tokens "
        "(confirmado permanente pela Anthropic em ago/2026 — não usar $3/$15)."
    )
    res["A14"] = (
        "•  Candidatas mais baratas (GPT-4o mini, Gemini Flash/Flash-Lite, ou Haiku como gerador) "
        "ficam na aba Comparativo Provedores — só entram em produção se empatar guardrails no bakeoff "
        "(ver docs/CUSTOS_LLM.md). Default do piloto permanece Anthropic Sonnet."
    )
    for r in range(10, 15):
        if res[f"A{r}"].value:
            res[f"A{r}"].alignment = wrap
            res.row_dimensions[r].height = 45

    comp = wb["Comparativo Provedores"]
    # Desfaz merges antigos (bullets mesclados) antes de reescrever
    for rng in list(comp.merged_cells.ranges):
        comp.unmerge_cells(str(rng))
    for r in range(4, 45):
        for c in range(1, 10):
            comp.cell(r, c).value = None

    comp["A1"] = "Comparativo — Anthropic vs OpenAI vs Gemini (tom Amigável)"
    comp["A2"] = (
        "Mesmas premissas de tokens da aba Premissas. Só o preço por token muda. "
        "Haiku como gerador NÃO é o stack atual (hoje Haiku = escopo). "
        "OpenAI e Gemini são candidatas ao bakeoff."
    )

    comp["A4"] = "Preços — candidatas (USD por milhão de tokens)"
    comp["A5"] = "Modelo"
    comp["B5"] = "Entrada"
    comp["C5"] = "Saída"
    comp["D5"] = "Nota"
    for col in ("A", "B", "C", "D"):
        comp[f"{col}5"].font = bold

    comp["A6"] = "Gemini 3.1 Flash-Lite (mais barato)"
    comp["B6"] = 0.25
    comp["C6"] = 1.5
    comp["D6"] = (
        "Gemini 2.5 Flash-Lite ($0,10/$0,40) mais barato, mas com aposentadoria anunciada — "
        "preferir 3.1 Flash-Lite como base barata Google."
    )
    comp["B6"].fill = yellow
    comp["C6"].fill = yellow

    comp["A7"] = "Gemini 3.6 Flash (faixa de qualidade)"
    comp["B7"] = 1.5
    comp["C7"] = 7.5
    comp["D7"] = (
        "Fonte: ai.google.dev/gemini-api/docs/pricing (consultado 12/08/2026). "
        "Usar só API paga — nunca free tier."
    )
    comp["B7"].fill = yellow
    comp["C7"].fill = yellow

    comp["A8"] = "OpenAI GPT-4o mini"
    comp["B8"] = 0.15
    comp["C8"] = 0.60
    comp["D8"] = "Fonte: openai.com/api/pricing. Já suportado via LLM_PROVIDER=openai no projeto."
    comp["B8"].fill = yellow
    comp["C8"].fill = yellow

    tin = "(Premissas!$B$14+Premissas!$B$15+Premissas!$B$16+Premissas!$B$17)"
    tout = "Premissas!$B$21"

    comp["A10"] = (
        "Custo por mensagem — tom Amigável (USD) — só geração "
        "(sem ponderar Haiku de escopo do stack atual)"
    )
    comp["A11"] = "Provedor / Modelo"
    comp["B11"] = "Custo por mensagem (USD)"
    comp["A11"].font = bold
    comp["B11"].font = bold

    comp["A12"] = "Claude Sonnet 5 (stack atual — geração)"
    comp["B12"] = f"={tin}/1000000*Premissas!$B$10+{tout}/1000000*Premissas!$C$10"

    comp["A13"] = (
        "Claude Haiku 4.5 (HIPÓTESE: usado como gerador — NÃO é o stack atual)"
    )
    comp["B13"] = f"={tin}/1000000*Premissas!$B$11+{tout}/1000000*Premissas!$C$11"

    comp["A14"] = "OpenAI GPT-4o mini (candidata bakeoff)"
    comp["B14"] = f"={tin}/1000000*$B$8+{tout}/1000000*$C$8"

    comp["A15"] = "Gemini 3.1 Flash-Lite (candidata bakeoff — barata)"
    comp["B15"] = f"={tin}/1000000*$B$6+{tout}/1000000*$C$6"

    comp["A16"] = "Gemini 3.6 Flash (candidata bakeoff — qualidade)"
    comp["B16"] = f"={tin}/1000000*$B$7+{tout}/1000000*$C$7"

    comp["A18"] = "Custo mensal por cenário (R$, tom Amigável)"
    comp["A19"] = "Cenário"
    comp["B19"] = "Mensagens/mês"
    comp["C19"] = "Sonnet 5"
    comp["D19"] = "Haiku como gerador*"
    comp["E19"] = "GPT-4o mini"
    comp["F19"] = "Gemini Flash-Lite"
    comp["G19"] = "Gemini 3.6 Flash"
    for col in range(1, 8):
        comp.cell(19, col).font = bold

    for i, src_row in enumerate([5, 6, 7], start=20):
        labels = {20: "Conservador", 21: "Moderado", 22: "Agressivo"}
        comp[f"A{i}"] = labels[i]
        comp[f"B{i}"] = f"=Cenários!$B${src_row}"
        comp[f"C{i}"] = f"=$B{i}*$B$12*Premissas!$B$6"
        comp[f"D{i}"] = f"=$B{i}*$B$13*Premissas!$B$6"
        comp[f"E{i}"] = f"=$B{i}*$B$14*Premissas!$B$6"
        comp[f"F{i}"] = f"=$B{i}*$B$15*Premissas!$B$6"
        comp[f"G{i}"] = f"=$B{i}*$B$16*Premissas!$B$6"

    comp["A23"] = (
        "* Haiku como gerador = hipótese de economia DENTRO da Anthropic; "
        "no código atual Haiku só classifica escopo."
    )
    comp["A23"].alignment = wrap

    comp["A25"] = "Leitura dos números"
    comp["A25"].font = bold
    comp["A26"] = (
        "•  No cenário Agressivo, Sonnet vs Gemini Flash-Lite fica na ordem de ~R$ 120–185/mês "
        '(não "poucas dezenas"). Isso é dinheiro real na linha de API — e ainda assim pequeno vs. 1 FTE. '
        "Dá para buscar a opção mais barata via bakeoff de qualidade, sem pressa."
    )
    comp["A27"] = (
        "•  Modelos baratos (Flash-Lite, às vezes mini) podem falhar em seguir instruções/guardrails. "
        "Para este bot (recusa segura, fail-closed, tom consistente), qualidade manda: "
        "barato só ganha se empatar o Sonnet nos testes."
    )
    comp["A28"] = (
        "•  Trocar de fornecedor reabre governança (dossiê SI / DPA). Trocar só o modelo dentro da "
        "Anthropic (ex.: Haiku como gerador) é mais leve em SI, mas ainda exige bakeoff."
    )
    comp["A29"] = (
        "•  IMPORTANTE: a camada gratuita do Gemini (Google AI Studio) pode reutilizar conteúdo "
        "para melhorar produtos Google — incompatível com o dossiê. Qualquer teste Gemini = "
        "API paga com faturamento ativado, nunca free tier."
    )
    comp["A29"].font = red_font
    comp["A30"] = (
        "•  Recomendação agora: manter Sonnet em produção/piloto; usar esta aba + docs/CUSTOS_LLM.md "
        "para experimentar candidatas baratas com a mesma bateria de testes. "
        "Só trocar default se empatar guardrails."
    )
    for r in range(26, 31):
        comp[f"A{r}"].alignment = wrap
        comp.row_dimensions[r].height = 52

    # Benefício esperado: ligar aos Cenários recalculados (não valores congelados $3/$15)
    ben = wb["Benefício esperado"]
    ben["B5"] = "=Cenários!B5"
    ben["B6"] = "=Cenários!B6"
    ben["B7"] = "=Cenários!B7"
    ben["B14"] = "=Cenários!F5"
    ben["B15"] = "=Cenários!F6"
    ben["B16"] = "=Cenários!F7"
    ben["A10"] = (
        "Compare o custo de API (aba Resumo/Cenários — dezenas a ~R$ 130/mês no Agressivo "
        "com Sonnet $2/$10) com o custo de oportunidade da equipe de RH/TI e com o ganho de "
        "orientação oficial na plataforma Enleva. Não multiplique mensagens por um "
        "“preço médio de chamado” inventado. Candidatas baratas: aba Comparativo + bakeoff."
    )
    ben["A10"].alignment = wrap

    # Textos de diferença alinhados ao preço vigente $2/$10 (~R$ 117 no Agressivo Sonnet vs Lite)
    res["A10"] = (
        "•  O custo de API no cenário Moderado fica ~R$ 50/mês (Sonnet $2/$10). No Agressivo, "
        "Sonnet vs Gemini Flash-Lite fica ~R$ 115–120/mês de diferença — material na linha de API, "
        "ainda pequeno vs 1 FTE. Há tempo para bakeoff de qualidade antes de trocar provedor."
    )
    comp["A26"] = (
        "•  No cenário Agressivo (preço Sonnet vigente $2/$10), Sonnet vs Gemini Flash-Lite fica "
        "~R$ 115–120/mês (não “poucas dezenas”). Dinheiro real na linha de API; ainda assim "
        "pequeno vs 1 FTE. Buscar a opção mais barata via bakeoff, sem pressa."
    )

    comp.column_dimensions["A"].width = 64
    comp.column_dimensions["B"].width = 16
    for col in ("C", "D", "E", "F", "G"):
        comp.column_dimensions[col].width = 16
    comp.column_dimensions["D"].width = 55

    wb.save(SRC)
    DST_DOCS.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DST_DOCS)

    # Expected figures (generation-only comparative + stack with scope)
    tin_v = 400 + 1200 + 30 + 300
    tout_v = 220
    fx = 5.3
    sonnet_msg = tin_v / 1e6 * 2 + tout_v / 1e6 * 10
    haiku_scope = (150 / 1e6 * 1 + 10 / 1e6 * 5) * 0.2
    total_amigavel = sonnet_msg + haiku_scope
    lite_msg = tin_v / 1e6 * 0.25 + tout_v / 1e6 * 1.5
    openai_msg = tin_v / 1e6 * 0.15 + tout_v / 1e6 * 0.60

    print("saved", SRC)
    print("saved", DST_DOCS)
    print("moderado stack Amigavel USD", round(1620 * total_amigavel, 4))
    print("moderado stack Amigavel BRL", round(1620 * total_amigavel * fx, 2))
    print("agressivo sonnet gen BRL", round(4200 * sonnet_msg * fx, 2))
    print("agressivo lite BRL", round(4200 * lite_msg * fx, 2))
    print("diff agressivo", round(4200 * (sonnet_msg - lite_msg) * fx, 2))
    print("moderado openai BRL", round(1620 * openai_msg * fx, 2))


if __name__ == "__main__":
    main()
